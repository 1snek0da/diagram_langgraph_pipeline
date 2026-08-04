"""Importer for user-authorized report exports; never scrapes protected terminals."""

from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path

from ..schemas.research import (
    CoverageReport,
    DatasetKind,
    EvidenceItem,
    FactBasis,
    MetricFact,
    SourceBatch,
    SourceDocument,
    SourceRequest,
    SourceTier,
)


class LicensedReportDirectoryAdapter:
    name = "licensed_report_directory"
    capabilities = frozenset(
        {
            DatasetKind.INDUSTRY_REPORTS,
            DatasetKind.CAPEX,
            DatasetKind.COMPANY_PROFILE,
            DatasetKind.PROFIT_FORECASTS,
            DatasetKind.MARGINAL_EVENTS,
        }
    )

    def __init__(self, roots: list[str | Path] | None = None):
        self.roots = tuple(Path(root).resolve() for root in (roots or []))

    def fetch(self, request: SourceRequest) -> SourceBatch:
        configured = [Path(value).resolve() for value in request.parameters.get("licensed_report_paths", [])]
        candidates = _collect_files([*self.roots, *configured])
        documents: list[SourceDocument] = []
        facts: list[MetricFact] = []
        evidence: list[EvidenceItem] = []
        data: dict = {}
        errors: list[str] = []
        warnings: list[str] = []
        for path in candidates:
            try:
                metadata = _read_sidecar(path)
                published_at = _metadata_datetime(metadata.get("published_at"))
                if published_at and published_at.date() > request.as_of_date:
                    continue
                content = path.read_bytes()
                text = _extract_text(path)
                document = SourceDocument(
                        provider=self.name,
                        source_tier=SourceTier.LICENSED_REPORT,
                        source_type=str(metadata.get("source_type", "licensed_report")),
                        external_id=metadata.get("external_id"),
                        title=str(metadata.get("title") or path.stem),
                        publisher=metadata.get("publisher"),
                        published_at=published_at,
                        file_path=str(path),
                        content_hash=sha256(content).hexdigest(),
                        language=str(metadata.get("language", "zh-CN")),
                        license_scope=str(metadata.get("license_scope", "authorized_local_use")),
                        text=text,
                        metadata={
                            key: value
                            for key, value in metadata.items()
                            if key not in {"published_at", "title", "facts", "evidence", "data"}
                        },
                    )
                documents.append(document)
                for item in metadata.get("facts", []):
                    facts.append(
                        MetricFact.model_validate(
                            {
                                **item,
                                "basis": item.get("basis", FactBasis.EXTRACTED),
                                "provider": self.name,
                                "source_id": document.source_id,
                            }
                        )
                    )
                for item in metadata.get("evidence", []):
                    evidence.append(
                        EvidenceItem.model_validate(
                            {
                                **item,
                                "source_id": document.source_id,
                                "extraction_method": item.get("extraction_method", "authorized_file_sidecar"),
                            }
                        )
                    )
                if isinstance(metadata.get("data"), dict):
                    data = _deep_merge(data, metadata["data"])
                if published_at is None:
                    warnings.append(f"{path.name} 缺少 sidecar published_at，不得用于 as_of_date 敏感预测")
            except Exception as exc:
                errors.append(f"{path}: {type(exc).__name__}: {exc}")
        requested = [str(path) for path in candidates]
        available = [document.file_path or document.title for document in documents]
        return SourceBatch(
            data=data,
            documents=documents,
            facts=facts,
            evidence=evidence,
            coverage=CoverageReport.from_items(requested, available, errors=errors),
            warnings=warnings,
            errors=errors,
        )


def _collect_files(roots: list[Path]) -> list[Path]:
    supported = {".pdf", ".docx", ".html", ".htm", ".txt"}
    result: list[Path] = []
    for root in roots:
        if root.is_file() and root.suffix.lower() in supported:
            result.append(root)
        elif root.is_dir():
            result.extend(path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in supported)
    return list(dict.fromkeys(result))


def _read_sidecar(path: Path) -> dict:
    sidecar = path.with_suffix(path.suffix + ".json")
    return json.loads(sidecar.read_text(encoding="utf-8")) if sidecar.exists() else {}


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = [
            " | ".join(cell.text.replace("\n", " / ") for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join([*paragraphs, *tables])
    if suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        return BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser").get_text("\n", strip=True)
    return path.read_text(encoding="utf-8", errors="replace")


def _metadata_datetime(value) -> datetime | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


def _deep_merge(left: dict, right: dict) -> dict:
    result = dict(left)
    for key, value in right.items():
        if isinstance(result.get(key), list) and isinstance(value, list):
            result[key] = [*result[key], *value]
        elif isinstance(result.get(key), dict) and isinstance(value, dict):
            result[key] = _deep_merge(result[key], value)
        else:
            result[key] = value
    return result
